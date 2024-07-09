from pypdf import PdfReader  # pdf
import ffmpeg  # audio/video
import xml.etree.ElementTree as ET  # xml
from PIL import Image, ExifTags
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
import datetime
import os
import subprocess
from pypdf.generic import NameObject, TextStringObject, DictionaryObject

OUTPUT_DIR = "output"

def extract_metadata_pdf(in_file) -> dict:
    reader = PdfReader(in_file)
    return reader.metadata


def extract_metadata_media(in_file) -> dict:
    try:
        # Run ffprobe
        metadata = ffmpeg.probe(in_file, show_format=None, show_streams=None)
        return metadata
    except FileNotFoundError as e:
        print("Error: ffprobe not found. Please ensure FFmpeg is installed and ffprobe is in your system's PATH.")
        raise e
    except ffmpeg.Error as e:
        print("An error occurred while extracting metadata from the media file:", e)
        print(e.stderr.decode('utf8'))
        raise e
    except Exception as e:
        print("An error occurred while extracting metadata from the media file:", e)
        raise e

def extract_metadata_raster_image(in_file) -> dict:
    image = Image.open(in_file)
    exif = image._getexif()
    if exif is None:
        return {}

    exif = {
        ExifTags.TAGS[k]: v
        for k, v in exif.items()
        if k in ExifTags.TAGS
    }

    # Convert Degree notation to decimal
    if "GPSInfo" in exif.keys():
        gps_tag = exif["GPSInfo"]
        N = gps_tag[2][0] + gps_tag[2][1] / 60.0 + gps_tag[2][2] / 3600.0
        E = gps_tag[4][0] + gps_tag[4][1] / 60.0 + gps_tag[4][2] / 3600.0
        exif["GPSInfo"] = f"{format(N, '.4f')} N, {format(E, '.4f')} E"  # Overwrite

    metadata = {k: v for k, v in exif.items() if v} if exif else {}
    return metadata

def extract_metadata_svg(in_file) -> dict:
    tree = ET.parse(in_file)
    root = tree.getroot()
    return root.attrib

def extract_metadata_xlsx(in_file) -> dict:
    workbook = load_workbook(in_file)
    properties = workbook.properties
    metadata = {
        'creator': properties.creator,
        'title': properties.title,
        'description': properties.description,
        'subject': properties.subject,
        'identifier': properties.identifier,
        'language': properties.language,
        'created': properties.created,
        'modified': properties.modified,
        'lastModifiedBy': properties.lastModifiedBy,
        'category': properties.category,
        'contentStatus': properties.contentStatus,
        'version': properties.version,
        'revision': properties.revision,
        'keywords': properties.keywords,
        'lastPrinted': properties.lastPrinted
    }
    return metadata

def extract_metadata_docx(in_file) -> dict:
    document = Document(in_file)
    properties = document.core_properties
    metadata = {
        'title': properties.title,
        'subject': properties.subject,
        'identifier': properties.identifier,
        'language': properties.language,
        'created': properties.created,
        'modified': properties.modified,
        'last_modified_by': properties.last_modified_by,
        'category': properties.category,
        'content_status': properties.content_status,
        'version': properties.version,
        'revision': properties.revision,
        'keywords': properties.keywords,
        'last_printed': properties.last_printed,
        'comments': properties.comments
    }
    return metadata

def extract_metadata_pptx(in_file) -> dict:
    presentation = Presentation(in_file)
    properties = presentation.core_properties
    metadata = {
        'title': properties.title,
        'subject': properties.subject,
        'identifier': properties.identifier,
        'language': properties.language,
        'created': properties.created,
        'modified': properties.modified,
        'last_modified_by': properties.last_modified_by,
        'category': properties.category,
        'content_status': properties.content_status,
        'version': properties.version,
        'revision': properties.revision,
        'keywords': properties.keywords,
        'last_printed': properties.last_printed,
        'comments': properties.comments
    }
    return metadata

def print_data(dictionary) -> None:
    if dictionary is None:
        print("No metadata found!")
        return
    items = dictionary.items()
    for k, v in items:
        print(k, ":", v)

def extract_metadata_file(file: str, file_type: str):
    metadata = {}

    match file_type:
        case "image/jpeg" | "image/png":
            metadata = extract_metadata_raster_image(file)
        case "application/pdf":
            metadata = extract_metadata_pdf(file)
        case "image/svg+xml":
            metadata = extract_metadata_svg(file)
        case "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
            metadata = extract_metadata_xlsx(file)
        case "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            metadata = extract_metadata_docx(file)
        case "application/vnd.openxmlformats-officedocument.presentationml.presentation":
            metadata = extract_metadata_pptx(file)
        case _:
            if "video/" in file_type or "audio/" in file_type:
                metadata = extract_metadata_media(file)
    return metadata

def add_meta_metadata(metadata: dict, file: str) -> dict:
    meta_metadata = DictionaryObject({
        NameObject('Metadata Creator'): TextStringObject('Zeta Data Extractor'),
        NameObject('Metadata Creation Date'): TextStringObject(datetime.datetime.now().isoformat()),
        NameObject('Metadata Last Modified By'): TextStringObject('Zeta Data Extractor'),
        NameObject('Metadata Last Modified Date'): TextStringObject(datetime.datetime.now().isoformat()),
        NameObject('Metadata Version'): TextStringObject('1.0'),
        NameObject('Metadata Status'): TextStringObject('Complete'),
        NameObject('Metadata Standard'): TextStringObject('Custom'),
        NameObject('Metadata Language'): TextStringObject('en'),
        NameObject('Metadata Source'): TextStringObject(file),
        NameObject('Metadata Accuracy'): TextStringObject('High'),
        NameObject('Metadata Completeness'): TextStringObject('Full'),
        NameObject('Metadata Coverage'): TextStringObject('All'),
        NameObject('Metadata Accessibility'): TextStringObject('Public'),
        NameObject('Metadata Rights'): TextStringObject('Open'),
        NameObject('Metadata Relationship'): TextStringObject('None')
    })
    metadata[NameObject('Meta-Metadata')] = meta_metadata
    return metadata

def create_result_folder():
    if not os.path.exists(OUTPUT_DIR):
        os.makedirs(OUTPUT_DIR)

def write_pdf(file: str, metadata: dict):
    # Implementation for writing the metadata back to a PDF
    pass
