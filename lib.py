from pypdf import PdfReader
from ffmpeg import FFmpeg
import xml.etree.ElementTree as ET
from PIL import Image, ExifTags
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
import datetime
import magic

def extract_metadata_pdf(in_file) -> dict:
    reader = PdfReader(in_file)
    metadata = {
        'Title': reader.metadata.get('/Title', ''),
        'Author': reader.metadata.get('/Author', ''),
        'Subject': reader.metadata.get('/Subject', ''),
        'Keywords': reader.metadata.get('/Keywords', ''),
        'Creator': reader.metadata.get('/Creator', ''),
        'Producer': reader.metadata.get('/Producer', ''),
        'Creation Date': reader.metadata.get('/CreationDate', ''),
        'Modification Date': reader.metadata.get('/ModDate', ''),
        'Trapped': reader.metadata.get('/Trapped', ''),
        'Page Count': len(reader.pages) if hasattr(reader, 'pages') else '',
        'Language': reader.metadata.get('/Lang', ''),
        'Encrypted': reader.is_encrypted if hasattr(reader, 'is_encrypted') else '',
        'PDF Version': reader.pdf_version if hasattr(reader, 'pdf_version') else '',
        'Tagged': reader.metadata.get('/Marked', '')
    }
    return {k: v for k, v in metadata.items() if v}

def extract_metadata_media(in_file) -> dict:
    ffprobe = FFmpeg(executable="ffprobe")
    ffprobe.input(in_file, print_format="xml", show_streams=None)
    metadata = ffprobe.execute()
    metadata_string = metadata.decode()
    tree = ET.fromstring(metadata_string)
    root = tree[0][0]
    return root.attrib

def extract_metadata_raster_image(in_file) -> dict:
    image = Image.open(in_file)
    exif = image._getexif()
    if exif:
        exif = {
            ExifTags.TAGS.get(k, k): v
            for k, v in exif.items()
        }
    return {k: v for k, v in exif.items() if v} if exif else {}

def extract_metadata_svg(in_file) -> dict:
    tree = ET.parse(in_file)
    root = tree.getroot()
    return root.attrib

def extract_metadata_xlsx(in_file) -> dict:
    workbook = load_workbook(in_file)
    properties = workbook.properties
    metadata = {
        'creator': properties.creator,
        'title': properties.title,
        'description': properties.description,
        'subject': properties.subject,
        'identifier': properties.identifier,
        'language': properties.language,
        'created': properties.created,
        'modified': properties.modified,
        'lastModifiedBy': properties.lastModifiedBy,
        'category': properties.category,
        'contentStatus': properties.contentStatus,
        'version': properties.version,
        'revision': properties.revision,
        'keywords': properties.keywords,
        'lastPrinted': properties.lastPrinted
    }
    return {k: v for k, v in metadata.items() if v}

def extract_metadata_docx(in_file) -> dict:
    document = Document(in_file)
    properties = document.core_properties
    metadata = {
        'Title': properties.title,
        'Subject': properties.subject,
        'Author': properties.author,
        'Keywords': properties.keywords,
        'Comments': properties.comments,
        'Last Modified By': properties.last_modified_by,
        'Creation Date': properties.created,
        'Last Modified Date': properties.modified,
        'Category': properties.category,
        'Content Status': properties.content_status,
        'Language': properties.language,
        'Version': properties.version,
        'Revision Number': properties.revision,
        'Total Editing Time': properties.total_editing_time,
        'Application Name': properties.application,
        'Template': properties.template
    }
    return {k: v for k, v in metadata.items() if v}

def extract_metadata_pptx(in_file) -> dict:
    presentation = Presentation(in_file)
    properties = presentation.core_properties
    metadata = {
        'Title': properties.title,
        'Subject': properties.subject,
        'Identifier': properties.identifier,
        'Language': properties.language,
        'Created': properties.created,
        'Modified': properties.modified,
        'Last Modified By': properties.last_modified_by,
        'Category': properties.category,
        'Content Status': properties.content_status,
        'Version': properties.version,
        'Revision': properties.revision,
        'Keywords': properties.keywords,
        'Last Printed': properties.last_printed,
        'Comments': properties.comments
    }
    return {k: v for k, v in metadata.items() if v}

def add_meta_metadata(metadata: dict, file: str) -> dict:
    metadata['Meta-Metadata'] = {
        'Metadata Creator': 'Zeta Data Extractor',
        'Metadata Creation Date': datetime.datetime.now().isoformat(),
        'Metadata Last Modified By': 'Zeta Data Extractor',
        'Metadata Last Modified Date': datetime.datetime.now().isoformat(),
        'Metadata Version': '1.0',
        'Metadata Status': 'Complete',
        'Metadata Standard': 'Custom',
        'Metadata Language': 'en',
        'Metadata Source': file,
        'Metadata Accuracy': 'High',
        'Metadata Completeness': 'Full',
        'Metadata Coverage': 'All',
        'Metadata Accessibility': 'Public',
        'Metadata Rights': 'Open',
        'Metadata Relationship': 'None'
    }
    return metadata

def print_data(dictionary) -> None:
    if not dictionary:
        print("No metadata found!")
        return
    for k, v in dictionary.items():
        if v:
            print(f"{k} : {v}")

def extract_all_metadata(file: str):
    file_type = magic.from_file(file, mime=True)
    metadata = {}
    
    match file_type:
        case "image/jpeg" | "image/png" | "image/tiff":
            metadata = extract_metadata_raster_image(file)
        case "application/pdf":
            metadata = extract_metadata_pdf(file)
        case "image/svg+xml":
            metadata = extract_metadata_svg(file)
        case "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
            metadata = extract_metadata_xlsx(file)
        case "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            metadata = extract_metadata_docx(file)
        case "application/vnd.openxmlformats-officedocument.presentationml.presentation":
            metadata = extract_metadata_pptx(file)
        case _:
            if "video/" in file_type or "audio/" in file_type:
                metadata = extract_metadata_media(file)

    metadata = add_meta_metadata(metadata, file)
    return metadata
